#!/usr/bin/env python3
"""Build a polished Word version of the TopoForesight TASE draft outline."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "244E6A"
DARK = "183447"
MUTED = "667085"
TABLE_FILL = "EEF2F5"
TABLE_BORDER = "AEB8C2"
CALLOUT_FILL = "F4F6F9"
PLACEHOLDER_FILL = "FFF2CC"
BODY_FONT = "Calibri"
EAST_ASIA_FONT = "SimSun"
MATH_FONT = "Cambria Math"
USABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=None, bold=None, italic=None, color=None,
                 font=BODY_FONT, east_asia=EAST_ASIA_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, bold, color):
    style.font.name = BODY_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    rfonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    rfonts.set(qn("w:cs"), BODY_FONT)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), TABLE_BORDER)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def column_widths(rows):
    cols = len(rows[0])
    scores = []
    for col in range(cols):
        values = []
        for row in rows[: min(len(rows), 12)]:
            text = row[col] if col < len(row) else ""
            length = sum(2 if "\u4e00" <= ch <= "\u9fff" else 1 for ch in text)
            values.append(min(max(length, 3), 30))
        scores.append(max(values) if values else 6)
    if cols >= 7:
        scores = [max(5, min(score, 16)) for score in scores]
    elif cols >= 5:
        scores = [max(6, min(score, 22)) for score in scores]
    else:
        scores = [max(8, score) for score in scores]
    total = sum(scores)
    widths = [max(650, int(USABLE_DXA * score / total)) for score in scores]
    scale = USABLE_DXA / sum(widths)
    widths = [int(width * scale) for width in widths]
    widths[-1] += USABLE_DXA - sum(widths)
    return widths


def configure_table(table, rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    set_table_borders(table)

    widths = column_widths(rows)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    font_size = 7.5 if len(widths) >= 7 else 8.0 if len(widths) >= 5 else 8.5
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            set_repeat_header(row)
        for col_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, TABLE_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                narrative = col_index == 0 or len(cell.text.strip()) > 18
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if narrative else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=(row_index == 0))


def shade_paragraph(paragraph, fill=CALLOUT_FILL, border_color=TABLE_BORDER):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), border_color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    visible = OxmlElement("w:t")
    visible.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, visible, end):
        run._r.append(node)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\\\(.+?\\\)|____|【待补引用：.+?】)")


def prettify_math(text):
    text = text.replace(r"\(", "").replace(r"\)", "")
    replacements = (
        (r"\rightarrow", "→"),
        (r"\cdots", "⋯"),
        (r"\ldots", "…"),
        (r"\geq", "≥"),
        (r"\leq", "≤"),
        (r"\neq", "≠"),
        (r"\in", "∈"),
        (r"\mid", "|"),
        (r"\Delta", "Δ"),
        (r"\delta", "δ"),
        (r"\lambda", "λ"),
        (r"\epsilon", "ε"),
        (r"\beta", "β"),
        (r"\gamma", "γ"),
        (r"\pi", "π"),
        (r"\star", "★"),
        (r"\times", "×"),
        (r"\mathbb E", "E"),
        (r"\begin{cases}", "{"),
        (r"\end{cases}", "}"),
        (r"\left", ""),
        (r"\right", ""),
        (r"\qquad", "   "),
        (r"\quad", "  "),
        (r"\\", " ; "),
    )
    for old, new in replacements:
        text = text.replace(old, new)
    for _ in range(4):
        text = re.sub(
            r"\\(?:mathbf|mathrm|mathcal|operatorname|text|mathrm)\{([^{}]*)\}",
            r"\1",
            text,
        )
    text = re.sub(r"\\hat\s*([A-Za-z])", lambda m: m.group(1) + "\u0302", text)
    text = re.sub(r"\\tilde\s*([A-Za-z])", lambda m: m.group(1) + "\u0303", text)
    text = re.sub(r"_\{([^{}]*)\}", r"_\1", text)
    text = re.sub(r"\^\{([^{}]*)\}", r"^\1", text)
    text = text.replace("{", "(").replace("}", ")")
    text = text.replace("&", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def clean_text(text):
    return text.replace(r"\(", "").replace(r"\)", "").replace(r"\_", "_")


def add_inline_runs(paragraph, text, default_size=None):
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(clean_text(text[cursor:match.start()]))
            set_run_font(run, size=default_size)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(clean_text(token[2:-2]))
            set_run_font(run, size=default_size, bold=True, color=DARK)
        elif token.startswith(r"\(") and token.endswith(r"\)"):
            run = paragraph.add_run(prettify_math(token))
            set_run_font(run, size=default_size, italic=True,
                         font=MATH_FONT, east_asia=MATH_FONT)
        elif token == "____":
            run = paragraph.add_run("____")
            set_run_font(run, size=default_size, bold=True, color="7A5A00")
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PLACEHOLDER_FILL)
            run._element.get_or_add_rPr().append(shd)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=default_size, italic=True, color="9B1C1C")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(clean_text(text[cursor:]))
        set_run_font(run, size=default_size)


def add_formula(doc, lines):
    text = " ".join(part.strip() for part in lines if part.strip())
    paragraph = doc.add_paragraph(style="Formula")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(prettify_math(text))
    set_run_font(run, size=10, font=MATH_FONT, east_asia=MATH_FONT)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells):
            rows.append(cells)
        index += 1
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows], index


def add_markdown_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            paragraph = table.cell(row_index, col_index).paragraphs[0]
            paragraph.clear()
            add_inline_runs(paragraph, value)
    configure_table(table, rows)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def set_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, 11, False, "111111")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.widow_control = True

    specs = {
        "Title": (24, True, DARK, 0, 10),
        "Subtitle": (12, False, MUTED, 0, 12),
        "Heading 1": (16, True, BLUE, 18, 10),
        "Heading 2": (13, True, BLUE, 12, 6),
        "Heading 3": (11.5, True, DARK, 8, 4),
        "Caption": (9, False, MUTED, 5, 6),
    }
    for name, (size, bold, color, before, after) in specs.items():
        style = doc.styles[name]
        set_style_font(style, size, bold, color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True

    formula = (
        doc.styles["Formula"]
        if "Formula" in doc.styles
        else doc.styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
    )
    set_style_font(formula, 10, False, "111111")
    formula.paragraph_format.space_before = Pt(5)
    formula.paragraph_format.space_after = Pt(7)
    formula.paragraph_format.keep_together = True

    for list_name in ("List Bullet", "List Number"):
        style = doc.styles[list_name]
        set_style_font(style, 11, False, "111111")
        style.paragraph_format.left_indent = Inches(0.38)
        style.paragraph_format.first_line_indent = Inches(-0.19)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208


def new_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        levels = abstract.findall(qn("w:lvl"))
        for level in levels:
            num_fmt = level.find(qn("w:numFmt"))
            if (
                level.get(qn("w:ilvl")) == "0"
                and num_fmt is not None
                and num_fmt.get(qn("w:val")) == "decimal"
            ):
                abstract_id = abstract.get(qn("w:abstractNumId"))
                break
        if abstract_id is not None:
            break
    if abstract_id is None:
        raise RuntimeError("No decimal numbering definition found")
    existing = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def apply_decimal_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_ref)


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("TASE-LEVEL JOURNAL DRAFT")
    set_run_font(run, size=10, bold=True, color=BLUE)
    kicker.paragraph_format.space_after = Pt(16)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True
    run = title.add_run("TopoForesight TASE 初稿骨架")
    set_run_font(run, size=24, bold=True, color=DARK)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "面向拓扑视觉语言导航的双阶段世界模型前瞻\n中文逻辑稿与实验占位版"
    )
    set_run_font(run, size=13, color=MUTED)

    for _ in range(2):
        doc.add_paragraph()

    details = [
        ("目标期刊", "IEEE Transactions on Automation Science and Engineering 水平上下"),
        ("主线依据", "docs/paper-mainline.md（2026-08-31）"),
        ("默认方法", "候选预算 K；前瞻展开步数 k=1"),
        ("数值状态", "全部实验数值保留为 ____，等待原始结果回填"),
        ("图像状态", "仅保留所需图片的详细图注，不生成图片"),
    ]
    table = doc.add_table(rows=len(details), cols=2)
    rows = []
    for row_index, (label, value) in enumerate(details):
        rows.append([label, value])
        table.cell(row_index, 0).text = label
        table.cell(row_index, 1).text = value
        set_cell_shading(table.cell(row_index, 0), TABLE_FILL)
    configure_table(table, rows)
    for row in table.rows:
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, size=9, bold=True, color=DARK)
        for run in row.cells[1].paragraphs[0].runs:
            set_run_font(run, size=9)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    shade_paragraph(note)
    run = note.add_run("本文档用于论文结构、方法公式和实验设计讨论；不是最终投稿排版稿。")
    set_run_font(run, size=10, italic=True, color=MUTED)
    doc.add_page_break()


def build(source_path, output_path):
    lines = source_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = True
    set_styles(doc)
    add_cover(doc)

    header = section.header.paragraphs[0]
    header.text = "TopoForesight · TASE 初稿骨架"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        set_run_font(run, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()

    doc.core_properties.title = "TopoForesight TASE 初稿骨架"
    doc.core_properties.subject = "TASE水平中文逻辑稿、公式、表格与图注占位"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "VLN-CE; TopoForesight; topological planning; world model; TASE"

    index = next(
        i for i, line in enumerate(lines)
        if line.startswith("## TASE样例反映出的结构原则")
    )
    pending_figure_box = False
    paper_title_seen = False
    numbered_group_active = False
    numbered_group_id = None
    while index < len(lines):
        text = lines[index].strip()
        if not text or text == "---":
            index += 1
            continue
        is_numbered_item = bool(re.match(r"^\d+\.\s+", text))
        if not is_numbered_item:
            numbered_group_active = False
            numbered_group_id = None
        if text.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            continue
        if text == r"\[":
            formula_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != r"\]":
                formula_lines.append(lines[index])
                index += 1
            add_formula(doc, formula_lines)
            index += 1
            continue
        if text.startswith("# "):
            if not paper_title_seen:
                doc.add_page_break()
                paper_title_seen = True
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            add_inline_runs(paragraph, text[2:], default_size=18)
            index += 1
            continue
        if text.startswith("#### "):
            paragraph = doc.add_paragraph(style="Heading 3")
            add_inline_runs(paragraph, text[5:])
            index += 1
            continue
        if text.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline_runs(paragraph, text[4:])
            index += 1
            continue
        if text.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline_runs(paragraph, text[3:])
            index += 1
            continue
        if is_numbered_item:
            if not numbered_group_active:
                numbered_group_id = new_decimal_numbering(doc)
                numbered_group_active = True
            paragraph = doc.add_paragraph(style="List Number")
            apply_decimal_numbering(paragraph, numbered_group_id)
            add_inline_runs(paragraph, re.sub(r"^\d+\.\s+", "", text))
            index += 1
            continue
        if text.startswith("- [ ] "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.16)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline_runs(paragraph, "☐ " + text[6:])
            index += 1
            continue
        if text.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, text[2:])
            index += 1
            continue
        if text.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.right_indent = Inches(0.08)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(7)
            shade_paragraph(paragraph)
            add_inline_runs(paragraph, text[2:])
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue
        if text.startswith("**图") and text.endswith("**"):
            paragraph = doc.add_paragraph(style="Caption")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
            add_inline_runs(paragraph, text)
            pending_figure_box = True
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline_runs(paragraph, text)
        if pending_figure_box:
            paragraph.paragraph_format.left_indent = Inches(0.12)
            paragraph.paragraph_format.right_indent = Inches(0.12)
            paragraph.paragraph_format.keep_together = True
            shade_paragraph(paragraph)
            for run in paragraph.runs:
                set_run_font(run, size=9.5, color=DARK)
            pending_figure_box = False
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_topoforesight_tase_docx.py <source.md> <output.docx>")
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())


if __name__ == "__main__":
    main()
