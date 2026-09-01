#!/usr/bin/env python3
"""Replace only the abstract paragraph in the current PILOT TASE draft."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path("/home/sia/project/ETP-R1/docs/TopoForesight_TASE_初稿骨架.docx")
OUTPUT = SOURCE

ABSTRACT = (
    "视觉语言导航（VLN）要求智能体理解自然语言指令，并在未见过的环境中进行连续感知和决策。"
    "在线拓扑方法能够利用历史观测组织长程导航，但对未访问候选的判断仍主要依赖当前视觉与拓扑记忆，"
    "难以获知候选方向之后可能出现的场景。为解决这一问题，本文提出PILOT，"
    "一种面向在线拓扑视觉语言导航的渐进式世界模型前瞻框架。PILOT利用冻结的表征世界模型，"
    "在动作执行前预测与拓扑候选对应的未来视觉信息。该框架首先为多个候选补充到达后的预测表征，"
    "再对少量关键候选进行进一步前瞻，并利用预测结果改善候选选择。与此同时，"
    "PILOT通过限制预测对基础策略的修正范围，并仅在可能影响当前决策时投入深入前瞻计算，"
    "提高方法的可靠性和计算效率。大量实验在R2R-CE和RxR-CE数据集上进行。结果表明，"
    "PILOT能够稳定改善视觉语言导航性能，并在导航效果与计算开销之间取得良好平衡。"
    "代码见：https://github.com/____/PILOT。"
)


def replace_paragraph_text(paragraph, text: str) -> None:
    """Replace visible content while retaining paragraph and first-run formatting."""
    first_rpr = None
    for run in paragraph.runs:
        if run._element.rPr is not None:
            first_rpr = deepcopy(run._element.rPr)
            break

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if first_rpr is not None:
        if run._element.rPr is not None:
            run._element.remove(run._element.rPr)
        run._element.insert(0, first_rpr)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    document = Document(SOURCE)
    abstract_index = None
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip() == "摘要":
            abstract_index = index
            break

    if abstract_index is None:
        raise RuntimeError("未找到摘要标题")

    body_index = abstract_index + 1
    while body_index < len(document.paragraphs) and not document.paragraphs[body_index].text.strip():
        body_index += 1
    if body_index >= len(document.paragraphs):
        raise RuntimeError("未找到摘要正文")

    replace_paragraph_text(document.paragraphs[body_index], ABSTRACT)
    document.save(OUTPUT)

    verify = Document(OUTPUT)
    paragraphs = [paragraph.text.strip() for paragraph in verify.paragraphs]
    if ABSTRACT not in paragraphs:
        raise RuntimeError("保存后的摘要校验失败")
    if not paragraphs or not paragraphs[0].startswith("PILOT:"):
        raise RuntimeError("保存后的标题校验失败")

    print(OUTPUT)


if __name__ == "__main__":
    main()
