#!/usr/bin/env python3
"""Apply the academic Chinese narrative-proposal style to the paper DOCX."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


EAST_ASIA_FONT = "SimSun"
LATIN_FONT = "Times New Roman"
HEADING_BLUE = "2E5D7B"
HEADING_DARK = "24485F"
MUTED = "6B7280"
TABLE_FILL = "F4F6F9"
TABLE_BORDER = "B8C2CC"
USABLE_DXA = 9411  # A4 width minus 2.2 cm margins on both sides.


def set_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = LATIN_FONT
    props = run._element.get_or_add_rPr()
    fonts = props.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    fonts.set(qn("w:cs"), LATIN_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, *, bold=False, color="000000"):
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    props = style.element.get_or_add_rPr()
    fonts = props.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    fonts.set(qn("w:cs"), LATIN_FONT)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), TABLE_BORDER)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def width_profile(header, columns):
    joined = "|".join(header)
    profiles = [
        ("项目|设置", [0.22, 0.78]),
        ("参数组|基础损失", [0.58, 0.21, 0.21]),
        ("维度|NavMorph", [0.18, 0.39, 0.43]),
        ("数据集|阶段|唯一新增因素", [0.11, 0.13, 0.31, 0.15, 0.15, 0.15]),
        ("变体|训练输入|评测输入|目的", [0.18, 0.25, 0.24, 0.33]),
        ("q1 来源|视野", [0.33, 0.15, 0.26, 0.26]),
        ("CLS 距离", [0.09, 0.18, 0.20, 0.18, 0.17, 0.18]),
        ("类别|数量|占全部比例", [0.34, 0.15, 0.17, 0.17, 0.17]),
        ("方法|NE", [0.36, 0.16, 0.16, 0.16, 0.16]),
        ("方法|SR|SPL|nDTW", [0.36, 0.16, 0.16, 0.16, 0.16]),
        ("数据集|阶段|SR|SPL", [0.18, 0.18, 0.16, 0.16, 0.16, 0.16]),
    ]
    for token, values in profiles:
        if token in joined and len(values) == columns:
            return values
    return [1.0 / columns] * columns


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    set_table_borders(table)

    cols = len(table.columns)
    if cols == 0 or not table.rows:
        return
    header = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
    fractions = width_profile(header, cols)
    widths = [int(USABLE_DXA * value) for value in fractions]
    widths[-1] += USABLE_DXA - sum(widths)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[column_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade_cell(cell, TABLE_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.10
                narrative = column_index == 0 or len(cell.text.strip()) > 26
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if narrative else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    set_font(run, size=8.5, bold=(row_index == 0))


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def style_document(path: Path):
    doc = Document(path)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    style_specs = {
        "Title": (18, True, "203B4A", 0, 8),
        "Subtitle": (11, False, MUTED, 0, 12),
        "Heading 1": (16, True, HEADING_BLUE, 18, 10),
        "Heading 2": (13, True, HEADING_BLUE, 12, 6),
        "Heading 3": (12, True, HEADING_DARK, 8, 4),
        "Caption": (9, False, MUTED, 4, 8),
        "Image Caption": (9, False, MUTED, 4, 8),
        "Bibliography": (9.5, False, "000000", 0, 3),
    }
    for name, (size, bold, color, before, after) in style_specs.items():
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        set_style_font(style, size, bold=bold, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
        if name in {"Caption", "Image Caption"}:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style.paragraph_format.keep_with_next = True
        if name == "Bibliography":
            style.paragraph_format.left_indent = Inches(0.25)
            style.paragraph_format.first_line_indent = Inches(-0.25)
            style.paragraph_format.line_spacing = 1.0

    for index, paragraph in enumerate(doc.paragraphs):
        style_name = paragraph.style.name if paragraph.style else ""
        text = paragraph.text.strip()
        if style_name == "Title" or index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(8)
            for run in paragraph.runs:
                set_font(run, size=18, bold=True, color="203B4A")
        elif style_name == "Subtitle" or text == "作者信息待定":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_font(run, size=10.5, color=MUTED)
        else:
            for run in paragraph.runs:
                set_font(run)
        if style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True

    for table in doc.tables:
        style_table(table)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "受约束表征世界模型主动前视 · 中文逻辑母稿"
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.space_after = Pt(0)
    for run in header_p.runs:
        set_font(run, size=8.5, color=MUTED)
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.clear()
    add_page_field(footer_p)

    doc.core_properties.title = "面向连续环境视觉语言导航的受约束表征世界模型主动前视"
    doc.core_properties.subject = "中文期刊论文逻辑母稿"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "VLN-CE; topology; world model; active lookahead"

    temporary = path.with_suffix(".styled.tmp.docx")
    doc.save(temporary)
    temporary.replace(path)


def main():
    if len(sys.argv) != 2:
        raise SystemExit("usage: style_paper_docx.py <paper.docx>")
    style_document(Path(sys.argv[1]).resolve())


if __name__ == "__main__":
    main()
